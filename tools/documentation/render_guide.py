#!/usr/bin/env python3
"""Render the source-backed PK DTS guide with editable PowerPoint callouts.

Screenshots remain original files embedded in the deck.  All crops use native
PowerPoint picture cropping; arrows, labels, and slide text remain editable.
Only standard library modules, Pillow, and python-pptx are required.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


NAVY = "17243B"
BURGUNDY = "A80037"
TEAL = "087F87"
INK = "293346"
MUTED = "657084"
PALE = "F2F5F8"
WHITE = "FFFFFF"
LINE = "DFE5ED"
GOLD = "EDA83D"
W, H = 13.333333, 7.5
FONT = "Calibri"
FONT_FILE = Path("C:/Windows/Fonts/calibri.ttf")

SCREENSHOTS = {
    1: "login", 2: "login", 3: "overview", 4: "hardcopy",
    5: "flow_hardcopy_create", 6: "flow_softcopy_create",
    7: "flow_softcopy_revise", 8: "draft_softcopy_create",
    9: "flow_access", 10: "flow_transfer", 11: "transfer_accept",
    12: "flow_disposal", 13: "approval_task", 14: "notifications",
    15: "direct_hardcopy", 16: "assign_user", 17: "places",
    18: "new_user", 19: "roles", 20: "workflow_steps",
    21: "audit", 22: "backups",
}
FALLBACKS = {
    "flow_hardcopy_create": "hardcopy", "flow_softcopy_create": "requests",
    "flow_softcopy_revise": "softcopy", "draft_softcopy_create": "requests",
    "flow_access": "requests", "flow_transfer": "requests",
    "transfer_accept": "tasks", "flow_disposal": "requests",
    "approval_task": "tasks", "notifications": "overview",
    "direct_hardcopy": "hardcopy", "assign_user": "assignments",
    "new_user": "users", "workflow_steps": "workflows",
    "audit": "overview", "backups": "overview",
}
FLOW_GUIDANCE = {
    "softcopy_create": 6, "softcopy_revise": 7, "softcopy_cancel": 7,
    "hardcopy_create": 5, "hardcopy_update": 5, "access": 9,
    "assignment": 9, "transfer": 10, "disposal": 12,
}
FLOW_LABELS = {
    "softcopy_create": "Create a softcopy", "softcopy_revise": "Revise a softcopy",
    "softcopy_cancel": "Cancel a softcopy", "hardcopy_create": "Create a hardcopy",
    "hardcopy_update": "Update a hardcopy", "access": "Request document access",
    "assignment": "Request document assignment", "transfer": "Transfer a hardcopy",
    "disposal": "Dispose of a hardcopy",
}
FIELD_LABELS = {
    "title": "Proposed Document Title", "subject": "Subject / Reason",
    "document_number": "Document Number", "series_number": "Series Number",
    "new_revision_level": "New Revision Level", "page_number": "Number of Pages",
    "effective_date": "Effective Date", "expires_at": "Access Expires At",
    "disposal_reason": "Disposal Reason", "disposal_other": "Other Reason",
    "category": "Category", "location": "Location", "recipient": "Recipient / Assigned User",
    "holder": "Current Holder", "file": "Document File", "remarks": "Remarks",
}


def color(value):
    return RGBColor.from_string(value)


def clean(value):
    return str(value or "").replace("\u2192", " → ").strip()


def title_text(value):
    return re.sub(r"^\d+\.\s*", "", clean(value))


def wrap_text(value, width_inches, font_size):
    """Wrap with the same installed font used by PowerPoint, at 4x precision."""
    if FONT_FILE.exists():
        font = ImageFont.truetype(str(FONT_FILE), max(1, round(font_size * 4)))
        measure = font.getlength
    else:
        measure = lambda s: len(s) * font_size * 4 * 0.50
    # A conservative inset avoids Office's slightly wider line-breaking metrics.
    max_width = width_inches * 72 * 4 * 0.88
    lines = []
    for paragraph in clean(value).split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = f"{line} {word}".strip()
            if line and measure(candidate) > max_width:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if not line else None
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(0.6)
    if radius:
        shape.adjustments[0] = 0.025 if h > 1 else 0.12
    return shape


def textbox(slide, x, y, w, h, value, size=14, fill=INK, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.10
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color(fill)
    return shape


def set_notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


class GuideDeck:
    def __init__(self, args):
        self.args = args
        self.model = read_json(args.model)
        self.sections = self.model.get("sections", [])
        self.by_number = {int(re.match(r"\d+", s["title"])[0]): s
                          for s in self.sections if re.match(r"\d+", s.get("title", ""))}
        self.screen_dir = Path(args.screenshots)
        annotation_path = Path(args.annotations) if args.annotations else self.screen_dir / "annotations.json"
        self.annotations = read_json(annotation_path) if annotation_path.exists() else {}
        if "screenshots" in self.annotations:
            self.annotations = self.annotations["screenshots"]
        if isinstance(self.annotations, list):
            self.annotations = {Path(v.get("file", v.get("screenshot", ""))).stem: v
                                for v in self.annotations}
        flow_data = read_json(args.flows) if args.flows else {}
        self.flows = flow_data.get("flows", flow_data.get("scenarios", [])) if isinstance(flow_data, dict) else flow_data
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.prs.core_properties.title = "PK Document Control — User Guide"
        self.prs.core_properties.subject = "Practical system modules and all nine demonstrated request flows"
        self.prs.core_properties.author = "PK Document Control"
        self.prs.core_properties.keywords = "user guide, screenshots, synthetic examples, document control"
        self.records, self.missing = [], []

    def blank(self, title, chapter="USER GUIDE", subtitle=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color(PALE)
        rect(slide, 0, 0, W, 0.09, BURGUNDY)
        textbox(slide, 0.38, 0.22, 11.8, 0.20, chapter.upper(), 9, BURGUNDY, True)
        title_size = 24 if len(title) < 64 else 21
        textbox(slide, 0.38, 0.48, 12.55, 0.45, title, title_size, NAVY, True)
        if subtitle:
            textbox(slide, 0.39, 0.96, 12.1, 0.29, subtitle, 10.5, MUTED)
        self.records.append({"title": title, "chapter": chapter})
        return slide

    def footer(self, slide, route="", sample=True):
        rect(slide, 0.38, 7.13, 12.57, 0.009, LINE)
        left = "PK DOCUMENT CONTROL  |  USER GUIDE"
        textbox(slide, 0.40, 7.24, 3.3, 0.15, left, 7.1, MUTED, True)
        middle = "Actual system screens · Synthetic demonstration data" if sample else "Source-backed interface instructions"
        textbox(slide, 3.5, 7.24, 6.4, 0.15, middle, 7.2, MUTED, align=PP_ALIGN.CENTER)
        textbox(slide, 11.8, 7.22, 1.12, 0.18, f"{len(self.prs.slides):02d}", 9, NAVY, True, PP_ALIGN.RIGHT)

    def resolve(self, stem):
        stem = clean(stem)
        name = Path(stem).name
        candidate = self.screen_dir / (name if Path(name).suffix else name + ".png")
        if candidate.exists():
            return candidate, Path(name).stem
        self.missing.append(str(candidate))
        if not self.args.allow_missing:
            return None, Path(name).stem
        fallback = FALLBACKS.get(Path(name).stem, "overview")
        candidate = self.screen_dir / (fallback + ".png")
        return (candidate if candidate.exists() else None), fallback

    def annotation_record(self, stem):
        return self.annotations.get(stem, self.annotations.get(stem + ".png", {})) or {}

    def screen_annotations(self, stem, local=None):
        values = local if local is not None else self.annotation_record(stem).get("annotations", [])
        if values:
            return values[:5]
        defaults = {
            "login": [(839, 529, 450, 48, "Username"), (839, 622, 450, 48, "Password"), (839, 689, 450, 48, "Sign In")],
            "overview": [(30, 323, 178, 35, "My Requests"), (30, 367, 178, 36, "My Tasks"), (1114, 17, 42, 42, "Notifications")],
            "hardcopy": [(1163, 142, 107, 41, "Table or Folder"), (317, 242, 896, 40, "Search records"), (1170, 406, 70, 35, "View details")],
            "softcopy": [(1163, 142, 107, 41, "Table or Folder"), (317, 242, 896, 40, "Search records")],
            "requests": [(1255, 195, 151, 42, "New Request")],
            "tasks": [(1253, 305, 144, 45, "Review assigned work")],
            "assignments": [(1230, 194, 175, 42, "Assign Document")],
            "places": [(1195, 195, 212, 42, "Add reference record")],
            "users": [(1275, 126, 111, 40, "Add User")],
            "roles": [(1275, 126, 111, 40, "Add Role")],
        }
        return [{"x": x, "y": y, "width": w, "height": h, "label": label}
                for x, y, w, h, label in defaults.get(stem, [])]

    def choose_crop(self, path, stem, annotations, requested=None):
        width, height = Image.open(path).size
        if requested:
            if isinstance(requested, dict):
                box = [requested.get("x", 0), requested.get("y", 0), requested.get("width", width), requested.get("height", height)]
            else:
                box = requested
            return self.bound_crop(box, width, height)
        if stem == "login":
            return [790, 135, 548, 694]
        if stem in ["overview", "notifications"]:
            return [0, 0, width, height]
        modal = stem.startswith("flow_") or stem in {
            "submit_confirmation", "approval_confirmation", "new_user", "assign_user",
            "workflow_steps", "direct_hardcopy", "direct_softcopy", "role_edit", "location_edit",
            "change_password", "request_returned", "request_rejected", "cancelled_draft",
            "transfer_dispatch", "transfer_accept",
        }
        if modal:
            # Dimmed backgrounds make the unmodified modal the bright component.
            with Image.open(path) as im:
                mask = im.convert("L").point(lambda p: 255 if p > 248 else 0)
                bounds = mask.getbbox()
            if (bounds and bounds[1] > 10 and bounds[2] < width - 10
                    and bounds[2] - bounds[0] > 280 and bounds[3] - bounds[1] > 150):
                x0, y0, x1, y1 = bounds
                return self.bound_crop([x0 - 28, y0 - 28, x1 - x0 + 56, y1 - y0 + 56], width, height)
        if stem.startswith("result_") or stem.startswith("draft_") or stem in {"approval_task", "return_previous_approver", "transfer_received"}:
            return self.bound_crop([278, 82, width - 290, 790], width, height)
        return self.bound_crop([278, 83, width - 290, height - 115], width, height)

    @staticmethod
    def bound_crop(box, width, height):
        x, y, w, h = [float(v) for v in box]
        x, y = max(0, x), max(0, y)
        return [x, y, max(1, min(w, width - x)), max(1, min(h, height - y))]

    def picture(self, slide, stem, annotations=None, crop=None):
        path, actual_stem = self.resolve(stem)
        panel = (0.38, 1.11, 9.10, 5.85)
        rect(slide, *panel, WHITE, LINE, True)
        if not path:
            textbox(slide, 0.8, 3.5, 8.2, 0.8, "Screenshot required: " + clean(stem), 18, BURGUNDY, True, PP_ALIGN.CENTER)
            return None
        screen_annotations = self.screen_annotations(actual_stem, annotations)
        requested_crop = crop or self.annotation_record(actual_stem).get("crop")
        cx, cy, cw, ch = self.choose_crop(path, actual_stem, screen_annotations, requested_crop)
        full_w, full_h = Image.open(path).size
        scale = min((panel[2] - 0.10) / cw, (panel[3] - 0.10) / ch)
        display_w, display_h = cw * scale, ch * scale
        left = panel[0] + (panel[2] - display_w) / 2
        top = panel[1] + (panel[3] - display_h) / 2
        pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(display_w), Inches(display_h))
        pic.crop_left, pic.crop_top = cx / full_w, cy / full_h
        pic.crop_right = max(0, 1 - (cx + cw) / full_w)
        pic.crop_bottom = max(0, 1 - (cy + ch) / full_h)
        pic.name = f"Original system screenshot: {path.name} (native picture crop)"
        visible = []
        for a in screen_annotations:
            x, y = float(a.get("x", 0)), float(a.get("y", 0))
            aw, ah = float(a.get("width", a.get("w", 1))), float(a.get("height", a.get("h", 1)))
            tx, ty = x + aw / 2, y + ah / 2
            if cx <= tx <= cx + cw and cy <= ty <= cy + ch:
                visible.append((a, tx, ty))
        # Labels occupy an inset rail.  Leaders point at DOM-derived controls.
        for i, (a, tx, ty) in enumerate(visible, 1):
            target_x, target_y = left + (tx - cx) * scale, top + (ty - cy) * scale
            side_left = tx < cx + cw * 0.50
            rail_x = panel[0] + 0.14 if side_left else panel[0] + panel[2] - 2.55
            rail_y = max(panel[1] + 0.12, min(panel[1] + panel[3] - 0.46, target_y - 0.50))
            if "label_y" in a:
                rail_y = panel[1] + float(a["label_y"])
            # Keep adjacent labels clear even when form fields share a row.
            for previous_y in [v[3] for v in getattr(self, "_callout_placements", []) if abs(v[0] - rail_x) < 0.3]:
                if abs(rail_y - previous_y) < 0.45:
                    rail_y = min(panel[1] + panel[3] - 0.46, previous_y + 0.47)
            if not hasattr(self, "_callout_placements"):
                self._callout_placements = []
            self._callout_placements.append((rail_x, target_x, target_y, rail_y))
            self.callout(slide, i, clean(a.get("label", "Control")), rail_x, rail_y, target_x, target_y)
        self._callout_placements = []
        self.records[-1]["screenshot"] = path.name
        self.records[-1]["crop"] = [round(v, 1) for v in (cx, cy, cw, ch)]
        self.records[-1]["annotation_count"] = len(visible)
        return path.name

    def callout(self, slide, number, label, x, y, target_x, target_y):
        # Straight leader and triangle arrowhead are editable OOXML geometry.
        start_x = x + (2.30 if target_x > x + 1.15 else 0.15)
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
            Inches(start_x), Inches(y + 0.18), Inches(target_x), Inches(target_y))
        connector.line.color.rgb = color(BURGUNDY)
        connector.line.width = Pt(2.2)
        line_xml = connector._element.spPr.get_or_add_ln()
        arrow = OxmlElement("a:tailEnd")
        arrow.set("type", "triangle")
        arrow.set("w", "med")
        arrow.set("len", "med")
        line_xml.append(arrow)
        rect(slide, x + 0.26, y, 2.15, 0.36, BURGUNDY, radius=True)
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y - 0.015), Inches(0.39), Inches(0.39))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color(WHITE)
        circle.line.color.rgb = color(BURGUNDY)
        circle.line.width = Pt(1.8)
        textbox(slide, x, y + 0.035, 0.39, 0.22, number, 11, BURGUNDY, True, PP_ALIGN.CENTER)
        label = label if len(label) <= 35 else label[:32].rstrip() + "…"
        textbox(slide, x + 0.44, y + 0.075, 1.88, 0.22, label, 8.4, WHITE, True)

    def narrative(self, slide, steps, notes=None, sample=None, badge=None):
        x, width, y = 9.75, 3.16, 1.14
        if badge:
            rect(slide, x, y, width, 0.34, TEAL, radius=True)
            textbox(slide, x + 0.12, y + 0.075, width - 0.24, 0.20, badge.upper(), 9.3, WHITE, True)
            y += 0.52
        if sample:
            textbox(slide, x, y, width, 0.22, "DEMONSTRATION INPUTS", 9.0, TEAL, True)
            y += 0.30
            for label, value in list(sample.items())[:7]:
                entry = f"{label}: {value}"
                wrapped = wrap_text(entry, width, 10.3)
                height = 0.185 * len(wrapped) + 0.06
                textbox(slide, x, y, width, height, "\n".join(wrapped), 10.3, INK)
                y += height + 0.04
            y += 0.10
        if steps:
            textbox(slide, x, y, width, 0.22, "WHAT TO DO", 9.0, BURGUNDY, True)
            y += 0.31
            note = " ".join(clean(n) for n in notes or [])
            if len(note) > 270:
                note = clean(notes[0])
            note_lines = wrap_text(note, width - 0.22, 9.2) if note else []
            note_height = max(0.55, len(note_lines) * 0.17 + 0.20) if note else 0
            room = 6.94 - y - (note_height + 0.12 if note else 0.05)
            sizes = [12.5, 12.0, 11.5, 11.0, 10.5]
            size = sizes[-1]
            for candidate in sizes:
                total_lines = sum(len(wrap_text(step, width - 0.30, candidate)) for step in steps)
                estimate = total_lines * candidate * 1.30 / 72 + len(steps) * 0.15
                if estimate < room:
                    size = candidate
                    break
            for i, step in enumerate(steps, 1):
                wrapped = wrap_text(step, width - 0.30, size)
                height = len(wrapped) * size * 1.30 / 72 + 0.035
                textbox(slide, x, y + 0.005, 0.23, 0.25, str(i).zfill(2), 9.5, TEAL, True)
                textbox(slide, x + 0.30, y, width - 0.30, height + 0.05, "\n".join(wrapped), size, INK)
                y += height + 0.12
        if notes:
            note = " ".join(clean(n) for n in notes)
            # Keep core caveats readable; full notes and sources are in speaker notes.
            if len(note) > 270:
                note = clean(notes[0])
            wrapped = wrap_text(note, width - 0.22, 9.2)
            estimated = max(0.55, len(wrapped) * 0.17 + 0.20)
            note_y = max(y + 0.02, 6.94 - estimated)
            if note_y + estimated <= 7.0:
                rect(slide, x, note_y, width, estimated, "E5F1F1", radius=True)
                textbox(slide, x + 0.11, note_y + 0.09, width - 0.22, estimated - 0.13, "\n".join(wrapped), 9.2, TEAL)

    def page(self, title, screenshot, steps, notes=None, route="", sources=None,
             chapter="USING THE SYSTEM", sample=None, badge=None, crop=None,
             annotations=None, full_notes=None):
        # Long procedures continue instead of silently losing steps.
        chunks = [steps[i:i + 5] for i in range(0, len(steps), 5)] or [[]]
        for i, chunk in enumerate(chunks):
            display_title = title if i == 0 else f"{title} — continued"
            slide = self.blank(display_title, chapter)
            self.picture(slide, screenshot, annotations, crop)
            self.narrative(slide, chunk, notes, sample if i == 0 else None, badge)
            self.footer(slide, route)
            detail = [f"Route: {route}", "Screenshot uses synthetic demonstration records.", "Steps:", *steps,
                      "Notes:", *(full_notes or notes or []), "Source references:", *(sources or [])]
            set_notes(slide, "\n".join(detail))
            self.records[-1]["route"] = route
            self.records[-1]["sources"] = sources or []

    def section(self, number):
        section = self.by_number.get(number)
        if not section:
            return
        steps = section.get("slide_steps", section.get("steps", []))
        self.page(title_text(section["title"]), section.get("screenshot", SCREENSHOTS[number]), steps,
                  section.get("slide_notes", section.get("notes", [])), section.get("route", ""),
                  section.get("sources", []), chapter=section.get("chapter", "SYSTEM MODULES"),
                  crop=section.get("crop"), annotations=section.get("annotations"),
                  full_notes=section.get("notes", []))

    def cover(self):
        slide = self.blank("", "")
        slide.shapes[0].fill.fore_color.rgb = color(TEAL)
        rect(slide, 0, 0.09, W, 7.41, NAVY)
        rect(slide, 0.58, 0.77, 2.52, 0.34, BURGUNDY, radius=True)
        textbox(slide, 0.74, 0.85, 2.2, 0.20, "PRACTICAL SYSTEM WALKTHROUGH", 8.6, WHITE, True)
        textbox(slide, 0.58, 1.45, 5.6, 1.50, "PK Document\nControl", 38, WHITE, True)
        textbox(slide, 0.60, 3.04, 5.2, 0.7, "USER GUIDE", 29, "F5B8C9", True)
        textbox(slide, 0.61, 3.97, 5.15, 1.16,
                "Find controlled records. Prepare requests.\nReview approvals. Complete physical handoffs.", 17, "D8E3F0")
        textbox(slide, 0.61, 5.37, 5.13, 0.88,
                "Actual application screens with editable arrows.\nNine request scenarios using synthetic sample data.", 12.5, "D8E3F0")
        path, _ = self.resolve("overview")
        if path:
            rect(slide, 6.15, 1.36, 6.6, 4.65, WHITE, radius=True)
            slide.shapes.add_picture(str(path), Inches(6.25), Inches(1.55), Inches(6.40), Inches(4.27))
        textbox(slide, 0.61, 6.89, 9, 0.23, "26 OCTOBER 2026  ·  Desktop application reference  ·  No real account credentials included", 9.6, "C7D5E8")
        self.records[-1]["title"] = "PK Document Control — User Guide"
        set_notes(slide, "Screenshots captured from the actual application using isolated, synthetic demonstration data. No generated passwords are included.")

    def roadmap(self):
        slide = self.blank("Your route through the workspace", "HOW TO USE THIS GUIDE")
        cards = [
            ("01", "GET ORIENTED", "Sign in and account security\nOverview and navigation\nFind records and inspect details"),
            ("02", "REQUEST AND REVIEW", "All nine request actions\nSave Draft → Submit → decision\nReturn, reject, dispatch, and accept"),
            ("03", "CONTROL AND ADMINISTER", "Direct documents and assignments\nPlaces, users, roles, and workflows\nAudit events and recovery tools"),
        ]
        for i, (n, title, body) in enumerate(cards):
            x = 0.43 + i * 4.23
            rect(slide, x, 1.58, 4.03, 3.76, WHITE, LINE, True)
            textbox(slide, x + 0.25, 1.92, 3.50, 0.65, n, 36, BURGUNDY, True)
            textbox(slide, x + 0.25, 2.86, 3.51, 0.6, title, 14.4, NAVY, True)
            textbox(slide, x + 0.25, 3.73, 3.49, 1.25, body, 15, MUTED)
        textbox(slide, 0.65, 5.83, 11.95, 0.72,
                "Numbered arrows identify the controls to use. Access and actions depend on your role, assigned documents, and the current workflow step.", 17, NAVY)
        textbox(slide, 0.65, 6.57, 11.95, 0.24, "All people, records, dates, and transactions shown as examples are synthetic demonstration data.", 11, MUTED)
        self.footer(slide)

    def flow(self, flow, index):
        typ = flow.get("type", "")
        guide = self.by_number.get(FLOW_GUIDANCE.get(typ), {})
        label = flow.get("label", flow.get("title", FLOW_LABELS.get(typ, typ.replace("_", " ").title())))
        sample = flow.get("sample", flow.get("sample_data", flow.get("data", flow.get("inputs", {}))))
        visible_keys = {
            "softcopy_create": {"title", "subject", "document_number", "series_number", "category", "file", "remarks"},
            "softcopy_revise": {"title", "subject", "new_revision_level", "page_number", "effective_date", "file", "remarks"},
            "softcopy_cancel": {"title", "subject", "remarks"},
            "hardcopy_create": {"title", "subject", "area", "specific", "asset", "location", "holder", "remarks"},
            "hardcopy_update": {"title", "subject", "area", "specific", "asset", "location", "holder", "remarks"},
            "access": {"title", "subject", "expires_at", "remarks"},
            "assignment": {"title", "subject", "remarks"},
            "transfer": {"title", "subject", "recipient", "location", "remarks"},
            "disposal": {"title", "subject", "disposal_reason", "disposal_other", "remarks"},
        }.get(typ, set(sample))
        sample = {FIELD_LABELS.get(k, k.replace("_", " ").title()): clean(v).title() if k == "disposal_reason" else clean(v)
                  for k, v in sample.items() if k in visible_keys and v not in (None, "")}
        if typ in ["hardcopy_create", "hardcopy_update"] and "Proposed Document Title" in sample:
            sample["Document Title"] = sample.pop("Proposed Document Title")
        elif typ == "softcopy_cancel" and "Proposed Document Title" in sample:
            sample["Existing Softcopy Document"] = sample.pop("Proposed Document Title")
        elif typ in ["transfer", "disposal"] and "Proposed Document Title" in sample:
            sample["Hardcopy Document"] = sample.pop("Proposed Document Title")
        elif typ in ["access", "assignment"] and "Proposed Document Title" in sample:
            sample["Selected document"] = sample.pop("Proposed Document Title")
        if typ == "transfer" and "Recipient / Assigned User" in sample:
            sample["Receiving User"] = sample.pop("Recipient / Assigned User")
        if typ in ["access", "assignment", "transfer"] and "Subject / Reason" in sample:
            sample["Request Subject" if typ == "transfer" else "Subject"] = sample.pop("Subject / Reason")
        section_steps = guide.get("flow_steps", {}).get(typ) or FLOW_STEPS.get(typ) or guide.get("slide_steps", guide.get("steps", []))[:4]
        route = flow.get("route", f"my-requests/{flow.get('tab', 'softcopy')}")
        self.page(f"Sample {index:02d} · {label}", flow.get("form", f"flow_{typ}"), section_steps,
                  ["Synthetic example. Save Draft first, then Submit after checking the proposal."], route,
                  guide.get("sources", []), "ALL NINE REQUEST FLOWS · INPUT", sample=sample,
                  badge="Worked example")
        status = clean(flow.get("status", flow.get("final_status", "")))
        result_stem = flow.get("result", f"result_{typ}")
        result_text = FLOW_RESULTS.get(typ, "Inspect the request status and its recorded result.")
        result_steps = [
            "Open the saved draft and verify the selected document, fields, and attachment where required.",
            "Select Submit, read the shared confirmation, and select Confirm. The request enters its published approval route.",
            "The assigned approver opens My Tasks, reviews the proposal, and selects Approve → Confirm at each required step.",
            result_text,
        ]
        if status:
            result_steps[-1] += f" The captured request status is {status.replace('_', ' ').title()}."
        self.page(f"Sample {index:02d} · {label} — result", result_stem, result_steps,
                  ["The screenshot shows the demonstrated result in the isolated sample workspace."], route,
                  guide.get("sources", []), "ALL NINE REQUEST FLOWS · OUTCOME", badge=f"Recorded result: {status or 'see screen'}")

    def extra(self, stem, title, steps, note, chapter="WORKED EXAMPLES", sources=None):
        # These demonstrations are optional for a baseline preview.  A supplied
        # flow model requires the complete result set at final generation.
        if not (self.screen_dir / (stem + ".png")).exists() and not self.flows:
            return
        self.page(title, stem, steps, [note], sources=sources, chapter=chapter)

    def appendix(self):
        unique = []
        for section in self.sections:
            for source in section.get("sources", []):
                if source not in unique:
                    unique.append(source)
        # All individual claims retain complete references in slide speaker notes.
        files = sorted(set(re.sub(r":\d+$", "", v) for v in unique))
        for chunk_index in range(0, len(files), 24):
            slide = self.blank("Interface evidence and source references", "REFERENCE APPENDIX")
            textbox(slide, 0.50, 1.20, 12.24, 0.52,
                    "Each module slide includes precise source paths and line references in its speaker notes. The guide follows the existing interface and verified demonstration flows.", 13, MUTED)
            chunk = files[chunk_index:chunk_index + 24]
            for j, item in enumerate(chunk):
                col, row = divmod(j, 12)
                textbox(slide, 0.55 + col * 6.34, 2.02 + row * 0.32, 6.08, 0.27,
                        item, 8.6, INK)
            textbox(slide, 0.55, 6.22, 12.2, 0.50,
                    "Screenshots: actual local application, 1440 × 960 capture. Demonstration records are synthetic. The original screenshot files are embedded without image editing; PowerPoint picture crops, arrows, and labels are editable.", 11, TEAL)
            self.footer(slide, sample=False)
            set_notes(slide, "\n".join(unique))

    def build(self):
        if self.flows and not self.args.allow_missing:
            types = [flow.get("type") for flow in self.flows]
            missing_types = sorted(set(FLOW_GUIDANCE) - set(types))
            if missing_types or len(types) != len(set(types)):
                raise ValueError("The final demonstration requires all nine unique request types. Missing: " + ", ".join(missing_types))
        self.cover()
        self.roadmap()
        if not self.flows:
            for section in self.sections:
                number = int(re.match(r"\d+", section["title"])[0])
                self.section(number)
        else:
            for n in [1, 2, 3, 4]:
                self.section(n)
            if (self.screen_dir / "softcopy_files.png").exists():
                self.extra("softcopy_files", "Open, download and print controlled files", [
                    "Open Softcopy Documents and find the permitted active record.",
                    "Select File History to inspect its available revision files.",
                    "Select Download for an authorized copy of the current approved file.",
                    "Select Print to open the controlled print preview."],
                    "Available file actions depend on your permissions and document access.",
                    sources=["application/models/File_model.php:7", "application/view/reusable_components/record_actions.php:63"])
            if (self.screen_dir / "file_history.png").exists():
                self.extra("file_history", "Inspect Document File History", [
                    "Open File History on the intended softcopy record.",
                    "Read the revision details and identify the current file.",
                    "Use available Download or Print actions for files within your access.",
                    "Select Close to return to the document register."],
                    "Historical revision metadata can remain after retired file bytes are cleaned up.",
                    sources=["application/view/pages/documents/shared_index.php:238", "application/models/File_model.php:7"])
            if (self.screen_dir / "print_file.png").exists():
                self.extra("print_file", "Use the controlled print preview", [
                    "Open Print from the authorized document file action.",
                    "Wait for the PDF, image or text preview to prepare.",
                    "Select Print when ready, then complete your browser's print dialog.",
                    "Use Download when you need a permitted local copy."],
                    "Office files use the downloaded file and an installed Office application for printing.",
                    sources=["application/view/pages/files/print.php:15", "application/controllers/Files.php:80"])
            self.section(8)
            self.extra("submit_confirmation", "Submit a checked draft", [
                "Save Draft records your proposal without starting approval.",
                "Back in the request list, select Submit for the intended draft.",
                "Read the request confirmation and select Confirm to start the approval workflow.",
                "Check the request becomes Submitted, then inspect its Approval Process Timeline."],
                "Submit uses the shared Confirm dialog; it does not ask for an approval remark.",
                sources=["public/assets/js/row-actions.js:19", "application/controllers/Requests.php:134"])
            for i, flow in enumerate(self.flows, 1):
                self.flow(flow, i)
            self.section(13)
            self.extra("approval_confirmation", "Approve the active workflow step", [
                "Open the assigned request in My Tasks and inspect the proposal and file.",
                "Select Approve for the active task.",
                "Read the shared confirmation and select Confirm.",
                "Check the next step or completed result in the Approval Process Timeline."],
                "Approve uses Confirm directly. It does not open a remark form.",
                sources=["public/assets/js/row-actions.js:19", "application/services/request/request_service.php:764"])
            self.extra("request_returned", "Return from the first approval step", [
                "At the first active step, select Return.",
                "Enter an optional Remark explaining what needs attention.",
                "Select Continue. This records the Return decision.",
                "The request becomes Returned to its requester, who may edit, save, and resubmit it."],
                "Continue in the Return dialog is the confirmation; no second Confirm dialog follows.",
                sources=["public/assets/js/row-actions.js:42", "application/services/request/request_service.php:792"])
            self.extra("return_previous_approver", "Return from a later approval step", [
                "Select Return when an earlier review needs to be reconsidered.",
                "Add an optional Remark and select Continue.",
                "The immediately previous approver becomes active again.",
                "After that approver approves again, the request returns to the step that returned it."],
                "A later-step Return keeps the request Submitted; it does not send an editable draft to the requester.",
                sources=["application/services/request/request_service.php:799"])
            self.extra("request_rejected", "Reject a proposal", [
                "Review the proposal and choose Reject when it should stop.",
                "Enter an optional Remark that explains the decision.",
                "Select Continue to record the rejection.",
                "Check Rejected in the request status and read the decision in the activity log."],
                "Rejected proposals are not completed document changes and are not editable returned drafts.",
                sources=["public/assets/js/row-actions.js:19", "application/services/request/request_service.php:820"])
            self.extra("cancelled_draft", "Cancel and, when eligible, remove a draft", [
                "Open your own Draft in My Requests and select Cancel.",
                "Confirm the cancellation and check the Cancelled status.",
                "If Delete is offered, check that this is the draft you intend to remove permanently.",
                "Select Delete and confirm. Only your cancelled, never-submitted draft is eligible."],
                "Permanent draft deletion cannot be undone. It is not an operation on an approved document.",
                sources=["application/services/request/request_policy.php:4"])
            self.section(11)
            self.extra("transfer_dispatch", "Record physical dispatch", [
                "After approval, open My Tasks → Hardcopy Transfer.",
                "The requester locates the intended Physical Transfer Handoff.",
                "Select Dispatch at physical release, then select Confirm.",
                "Check that the handoff awaits the receiving user's acceptance."],
                "Only the transfer requester can dispatch.", sources=["application/controllers/Requests.php:169"])
            self.extra("transfer_received", "Verify the received hardcopy", [
                "The receiving user opens the in-transit handoff.",
                "After checking the physical record, select Accept → Confirm.",
                "Inspect the updated holder and location in Hardcopy Documents.",
                "Use Transfer History when confirming the recorded handoff."],
                "The receiving user's acceptance completes physical receipt and updates custody.",
                sources=["application/controllers/Requests.php:184", "application/models/Document_detail_model.php:223"])
            self.section(14)
            self.section(15)
            if (self.screen_dir / "direct_softcopy.png").exists():
                self.extra("direct_softcopy", "Apply a direct softcopy action", [
                    "Open Softcopy Documents and select Softcopy Direct.",
                    "Choose Create Softcopy, Revise Softcopy, or Cancel Softcopy.",
                    "Complete the action-specific fields and attach a file for Create or Revise.",
                    "Select Approve Directly, then Confirm. Inspect the resulting record."],
                    "Direct actions require the appropriate permission and apply the change immediately.",
                    sources=["application/controllers/Documents.php:125", "application/view/pages/documents/shared_index.php:227"])
            for n in [16, 17, 18, 19, 20, 21, 22]:
                self.section(n)
        self.appendix()
        if self.missing and not self.args.allow_missing:
            raise ValueError("Required screenshots are missing:\n" + "\n".join(sorted(set(self.missing))))
        output = Path(self.args.output)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(output)
        manifest = output.with_suffix(".slides.json")
        manifest.write_text(json.dumps({"slides": self.records, "missing_screenshots": sorted(set(self.missing))}, indent=2), encoding="utf-8")
        print(json.dumps({"output": str(output), "slides": len(self.prs.slides), "manifest": str(manifest),
                          "missing_screenshots": sorted(set(self.missing))}, indent=2))


FLOW_STEPS = {
    "softcopy_create": ["Choose Create Softcopy and enter Subject / Reason.", "Enter the proposed title, Document Number, optional Series Number, and Category.", "Attach the Document File, add Remarks if needed, and select Save Draft."],
    "softcopy_revise": ["Choose Revise Softcopy and select the Existing Softcopy Document.", "Check the proposed title; enter New Revision Level and Number of Pages.", "Choose an optional Effective Date, attach the revised file, then Save Draft."],
    "softcopy_cancel": ["Choose Cancel Softcopy and select the Existing Softcopy Document.", "Enter Subject / Reason and any Remarks explaining the cancellation.", "Save Draft, check the selected record, and submit the proposal."],
    "hardcopy_create": ["Choose Create and enter Subject / Reason and Document Title.", "Select Area, Specific, optional Asset, Location, and check Current Holder.", "Optionally enable retention monitoring, then enter both Start and End dates (End on/after Start). Add notes and Save Draft."],
    "hardcopy_update": ["Choose Update and select the Existing Hardcopy Document.", "Check Document Title, Area, Specific, optional Asset, Location, and Current Holder.", "Update the intended metadata or retention information, then Save Draft."],
    "access": ["Choose Document Access Grants and enter Subject.", "Choose Document Type, select the document, and check your read-only Recipient field.", "Set Access Expires At, add optional Remarks, then Save Draft."],
    "assignment": ["Choose Document Assignments and enter Subject.", "Choose Document Type and select the document to be assigned to you.", "Check your read-only Recipient / Assigned User, add optional Remarks, then Save Draft."],
    "transfer": ["Select the Hardcopy Document and check Original Places and Holder.", "Choose the destination hierarchy and Receiving User; destination must differ from the original.", "Enter Request Subject and optional Remarks, then Save Draft."],
    "disposal": ["Choose Hardcopy Disposal, enter Subject / Reason, and select the document.", "Choose Disposal Reason: Shred, Scratch, or Other. Complete Other Reason if selected.", "Add optional Remarks, Save Draft, then review and submit the proposal."],
}
FLOW_RESULTS = {
    "softcopy_create": "The new controlled softcopy appears in Softcopy Documents after final approval.",
    "softcopy_revise": "The approved revision becomes the current controlled file; check the revision and file history.",
    "softcopy_cancel": "The controlled softcopy is cancelled after final approval; use the Cancelled status filter to inspect it.",
    "hardcopy_create": "The new hardcopy record appears with its approved storage hierarchy and holder.",
    "hardcopy_update": "The approved hardcopy metadata is applied; inspect the changed storage or retention fields.",
    "access": "The approved access grant applies to the selected document and the signed-in recipient until its expiry.",
    "assignment": "The approved assignment applies to the selected document and signed-in recipient.",
    "transfer": "Approval creates the handoff work. Complete Dispatch and recipient Accept before treating the physical transfer as received.",
    "disposal": "The record becomes Disposed; inspect Disposal Details and the selected reason in the document register.",
}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model", required=True, help="Guide model JSON with sections")
    parser.add_argument("--screenshots", required=True, help="Directory of original captured PNGs")
    parser.add_argument("--output", required=True, help="Destination .pptx file")
    parser.add_argument("--flows", help="Verified synthetic demonstration flows JSON")
    parser.add_argument("--annotations", help="DOM bounding-box manifest; defaults to screenshots/annotations.json")
    parser.add_argument("--allow-missing", action="store_true", help="Use baseline fallback images for an explicitly unfinished preview")
    GuideDeck(parser.parse_args()).build()


if __name__ == "__main__":
    main()
